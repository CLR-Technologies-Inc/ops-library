"""Build the Excel / LibreOffice Calc templates from the sample data.

    python build_workbooks.py

Every result cell is a live formula, so the workbooks recalculate when you
paste in your own data. Only functions that exist in both Excel (2010+) and
LibreOffice Calc are used. Requires openpyxl.
"""

import csv
import os
import sys

from openpyxl import Workbook as _Workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "data")
sys.path.insert(0, os.path.join(HERE, "..", "python"))
FMEA_DIR = os.path.join(HERE, "..", "..", "..", "03-design-and-engineering", "tools", "excel")

TITLE = Font(bold=True, size=14, color="1F4E79")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
RESULT_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NOTE = Font(italic=True, color="595959")


def Workbook():
    """A workbook that recalculates every formula when it is opened."""
    wb = _Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    return wb


def read(name):
    with open(os.path.join(DATA, name)) as f:
        return list(csv.DictReader(f))


def header(ws, row, col, labels):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def label_value(ws, row, label, formula, fmt="0.0000", fill=RESULT_FILL, col=1):
    ws.cell(row=row, column=col, value=label).font = BOLD
    c = ws.cell(row=row, column=col + 1, value=formula)
    c.number_format, c.fill, c.border = fmt, fill, BOX
    return c


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def line_chart(title_text, ws, cols, min_row, max_row, x_col, anchor, y_title):
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = title_text, y_title, "Sample"
    ch.height, ch.width = 9, 22
    for col in cols:
        ch.add_data(Reference(ws, min_col=col, min_row=min_row - 1, max_row=max_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=x_col, min_row=min_row, max_row=max_row))
    styles = ["1F4E79", "548235", "C00000", "C00000"]
    for s, color in zip(ch.series, styles):
        s.graphicalProperties.line.solidFill = color
        s.smooth = False
    for s in ch.series[2:]:
        s.graphicalProperties.line.dashStyle = "dash"
    ch.series[0].marker.symbol = "circle"
    ws.add_chart(ch, anchor)


# --------------------------------------------------------------- capability

def capability_workbook():
    values = [float(r["diameter_mm"]) for r in read("capability.csv")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Capability"
    title(ws, "Process Capability Study",
          "Yellow cells are inputs. Paste up to 500 individual values in column B "
          "(in time order). Green cells are results.")
    ws["A4"], ws["A5"], ws["A6"] = "LSL", "USL", "Target (optional)"
    for r, v in ((4, 9.9), (5, 10.1), (6, 10.0)):
        ws.cell(row=r, column=2, value=v).fill = INPUT_FILL
        ws.cell(row=r, column=1).font = BOLD
    first, last = 10, 509
    header(ws, 9, 1, ["Sample", "Value", "Moving range"])
    for i in range(first, last + 1):
        ws.cell(row=i, column=1, value=i - first + 1)
        v = values[i - first] if i - first < len(values) else None
        ws.cell(row=i, column=2, value=v).fill = INPUT_FILL
        if i > first:
            ws.cell(row=i, column=3,
                    value='=IF(AND(ISNUMBER(B%d),ISNUMBER(B%d)),ABS(B%d-B%d),"")' % (i, i - 1, i, i - 1))
    rng, mr = "B%d:B%d" % (first, last), "C%d:C%d" % (first + 1, last)
    col = 5
    rows = [
        ("n", "=COUNT(%s)" % rng, "0"),
        ("Mean", "=AVERAGE(%s)" % rng, "0.0000"),
        ("Std dev overall (long term)", "=STDEV(%s)" % rng, "0.000000"),
        ("Average moving range", "=AVERAGE(%s)" % mr, "0.000000"),
        ("Std dev within (MR-bar / 1.128)", "=F7/1.128379", "0.000000"),
        ("Cp", "=(B5-B4)/(6*F8)", "0.000"),
        ("Cpk", "=MIN(B5-F5,F5-B4)/(3*F8)", "0.000"),
        ("Cpm", '=IF(ISNUMBER(B6),(B5-B4)/(6*SQRT(F8^2+(F5-B6)^2)),"")', "0.000"),
        ("Pp", "=(B5-B4)/(6*F6)", "0.000"),
        ("Ppk", "=MIN(B5-F5,F5-B4)/(3*F6)", "0.000"),
        ("Expected ppm < LSL (overall)", "=NORMSDIST((B4-F5)/F6)*1000000", "0.0"),
        ("Expected ppm > USL (overall)", "=(1-NORMSDIST((B5-F5)/F6))*1000000", "0.0"),
        ("Expected ppm total (overall)", "=F14+F15", "0.0"),
        ("Observed ppm", "=(COUNTIF(%s,\"<\"&B4)+COUNTIF(%s,\">\"&B5))/F4*1000000" % (rng, rng), "0.0"),
        ("Sigma level (with 1.5 shift)", "=NORMSINV(1-F16/1000000)+1.5", "0.00"),
    ]
    header(ws, 3, col, ["Statistic", "Result"])
    for i, (lab, f, fmt) in enumerate(rows):
        label_value(ws, 4 + i, lab, f, fmt, col=col)
    ws.cell(row=20, column=col, value="Interpretation: Cpk of 1.33 or more is the usual minimum for an "
            "existing process; 1.67 for a new or safety-critical one. A big gap between Cpk and Ppk "
            "means the process drifts over time.").font = NOTE
    widths(ws, {"A": 18, "B": 12, "C": 14, "D": 3, "E": 34, "F": 14})

    # I-MR chart sheet
    ch = wb.create_sheet("I-MR Chart")
    title(ch, "Individuals and Moving Range Chart", "Reads the values on the Capability sheet. "
          "The chart shows the first 100 samples; extend the chart ranges for more.")
    header(ch, 4, 1, ["Sample", "Value", "CL", "UCL", "LCL", "MR", "MR-bar", "MR UCL", "Rule 1"])
    for i in range(500):
        r = 5 + i
        src = first + i
        ch.cell(row=r, column=1, value=i + 1)
        ch.cell(row=r, column=2, value='=IF(ISNUMBER(Capability!B%d),Capability!B%d,NA())' % (src, src))
        ch.cell(row=r, column=3, value="=Capability!$F$5")
        ch.cell(row=r, column=4, value="=Capability!$F$5+3*Capability!$F$8")
        ch.cell(row=r, column=5, value="=Capability!$F$5-3*Capability!$F$8")
        ch.cell(row=r, column=6, value='=IF(ISNUMBER(Capability!C%d),Capability!C%d,NA())' % (src, src))
        ch.cell(row=r, column=7, value="=Capability!$F$7")
        ch.cell(row=r, column=8, value="=(1+3*0.8525/1.128379)*Capability!$F$7")
        ch.cell(row=r, column=9, value='=IF(ISNUMBER(B%d),IF(OR(B%d>D%d,B%d<E%d),"OUT",""),"")'
                % (r, r, r, r, r))
    line_chart("Individuals", ch, [2, 3, 4, 5], 5, 104, 1, "K4", "Value")
    line_chart("Moving range", ch, [6, 7, 8], 5, 104, 1, "K24", "Moving range")
    widths(ch, {c: 10 for c in "ABCDEFGHI"})
    wb.save(os.path.join(HERE, "capability-and-imr.xlsx"))


# ------------------------------------------------------------------ X-bar R

def xbar_r_workbook():
    groups = {}
    for r in read("subgroups.csv"):
        groups.setdefault(int(r["subgroup"]), []).append(float(r["diameter_mm"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Xbar-R"
    title(ws, "X-bar and R Chart", "Enter up to 50 subgroups of equal size (2 to 10) in columns B to K. "
          "The subgroup size is taken from row 6.")
    header(ws, 5, 1, ["Subgroup"] + ["x%d" % i for i in range(1, 11)] +
           ["n", "Mean", "Range", "X-dbar", "UCL X", "LCL X", "R-bar", "UCL R", "LCL R", "Signal"])
    first, last = 6, 55
    for i in range(first, last + 1):
        g = i - first + 1
        ws.cell(row=i, column=1, value=g)
        vals = groups.get(g, [])
        for j in range(10):
            ws.cell(row=i, column=2 + j, value=vals[j] if j < len(vals) else None).fill = INPUT_FILL
        rowr = "B%d:K%d" % (i, i)
        ws.cell(row=i, column=12, value="=COUNT(%s)" % rowr)
        ws.cell(row=i, column=13, value='=IF(L%d>0,AVERAGE(%s),NA())' % (i, rowr))
        ws.cell(row=i, column=14, value='=IF(L%d>0,MAX(%s)-MIN(%s),NA())' % (i, rowr, rowr))
        ws.cell(row=i, column=15, value="=$C$3")
        ws.cell(row=i, column=16, value="=$C$3+$E$3*$D$3")
        ws.cell(row=i, column=17, value="=$C$3-$E$3*$D$3")
        ws.cell(row=i, column=18, value="=$D$3")
        ws.cell(row=i, column=19, value="=$G$3*$D$3")
        ws.cell(row=i, column=20, value="=$F$3*$D$3")
        ws.cell(row=i, column=21, value='=IF(L%d>0,IF(OR(M%d>P%d,M%d<Q%d,N%d>S%d),"OUT",""),"")'
                % (i, i, i, i, i, i, i))
        for col in range(13, 21):
            ws.cell(row=i, column=col).number_format = "0.0000"
    header(ws, 2, 2, ["n", "X-dbar", "R-bar", "A2", "D3", "D4", "d2", "Sigma within"])
    ws["B3"] = "=L6"
    ws["C3"] = '=AVERAGEIF(L6:L55,">0",M6:M55)'
    ws["D3"] = '=AVERAGEIF(L6:L55,">0",N6:N55)'
    ws["E3"] = "=VLOOKUP($B$3,Constants!$A$2:$H$10,2,FALSE)"
    ws["F3"] = "=VLOOKUP($B$3,Constants!$A$2:$H$10,3,FALSE)"
    ws["G3"] = "=VLOOKUP($B$3,Constants!$A$2:$H$10,4,FALSE)"
    ws["H3"] = "=VLOOKUP($B$3,Constants!$A$2:$H$10,5,FALSE)"
    ws["I3"] = "=D3/H3"
    for c in "BCDEFGHI":
        ws[c + "3"].fill, ws[c + "3"].number_format = RESULT_FILL, "0.0000"
    line_chart("X-bar chart", ws, [13, 15, 16, 17], first, first + 24, 1, "W2", "Subgroup mean")
    line_chart("Range chart", ws, [14, 18, 19, 20], first, first + 24, 1, "W22", "Range")
    widths(ws, {"A": 9})

    k = wb.create_sheet("Constants")
    header(k, 1, 1, ["n", "A2", "D3", "D4", "d2", "A3", "B3", "B4"])
    from opex_stats.constants import factors
    for n in range(2, 11):
        f = factors(n)
        for j, key in enumerate(["n", "A2", "D3", "D4", "d2", "A3", "B3", "B4"]):
            c = k.cell(row=n, column=1 + j, value=f[key] if key != "n" else n)
            c.number_format = "0" if key == "n" else "0.0000"
    wb.save(os.path.join(HERE, "xbar-r-chart.xlsx"))


# --------------------------------------------------------- attribute charts

def attribute_workbook():
    rows = read("attribute.csv")
    wb = Workbook()
    p = wb.active
    p.title = "p chart"
    title(p, "p Chart (proportion defective)", "Enter the number inspected and the number defective "
          "for up to 100 samples. Limits adjust to each sample size.")
    header(p, 4, 1, ["Sample", "Inspected", "Defective", "p", "p-bar", "UCL", "LCL", "Signal"])
    for i in range(100):
        r = 5 + i
        d = rows[i] if i < len(rows) else None
        p.cell(row=r, column=1, value=i + 1)
        p.cell(row=r, column=2, value=int(d["inspected"]) if d else None).fill = INPUT_FILL
        p.cell(row=r, column=3, value=int(d["defective"]) if d else None).fill = INPUT_FILL
        p.cell(row=r, column=4, value='=IF(ISNUMBER(B%d),C%d/B%d,NA())' % (r, r, r))
        p.cell(row=r, column=5, value="=$B$3")
        p.cell(row=r, column=6, value='=IF(ISNUMBER(B%d),MIN(1,$B$3+3*SQRT($B$3*(1-$B$3)/B%d)),NA())' % (r, r))
        p.cell(row=r, column=7, value='=IF(ISNUMBER(B%d),MAX(0,$B$3-3*SQRT($B$3*(1-$B$3)/B%d)),NA())' % (r, r))
        p.cell(row=r, column=8, value='=IF(ISNUMBER(B%d),IF(OR(D%d>F%d,D%d<G%d),"OUT",""),"")' % (r, r, r, r, r))
    label_value(p, 3, "p-bar", "=SUM(C5:C104)/SUM(B5:B104)", "0.00000")
    line_chart("p chart", p, [4, 5, 6, 7], 5, 4 + len(rows), 1, "J4", "Proportion defective")

    u = wb.create_sheet("u chart")
    title(u, "u Chart (defects per unit)", "Enter units inspected and defects found for up to 100 samples.")
    header(u, 4, 1, ["Sample", "Units", "Defects", "u", "u-bar", "UCL", "LCL", "Signal"])
    for i in range(100):
        r = 5 + i
        d = rows[i] if i < len(rows) else None
        u.cell(row=r, column=1, value=i + 1)
        u.cell(row=r, column=2, value=int(d["units"]) if d else None).fill = INPUT_FILL
        u.cell(row=r, column=3, value=int(d["defects"]) if d else None).fill = INPUT_FILL
        u.cell(row=r, column=4, value='=IF(ISNUMBER(B%d),C%d/B%d,NA())' % (r, r, r))
        u.cell(row=r, column=5, value="=$B$3")
        u.cell(row=r, column=6, value='=IF(ISNUMBER(B%d),$B$3+3*SQRT($B$3/B%d),NA())' % (r, r))
        u.cell(row=r, column=7, value='=IF(ISNUMBER(B%d),MAX(0,$B$3-3*SQRT($B$3/B%d)),NA())' % (r, r))
        u.cell(row=r, column=8, value='=IF(ISNUMBER(B%d),IF(OR(D%d>F%d,D%d<G%d),"OUT",""),"")' % (r, r, r, r, r))
    label_value(u, 3, "u-bar", "=SUM(C5:C104)/SUM(B5:B104)", "0.00000")
    line_chart("u chart", u, [4, 5, 6, 7], 5, 4 + len(rows), 1, "J4", "Defects per unit")

    c = wb.create_sheet("c and np")
    title(c, "c Chart and np Chart (constant sample size)", "c: defects per sample. np: defective units "
          "per sample, with the same sample size every time.")
    label_value(c, 3, "np sample size", 200, "0", INPUT_FILL)
    header(c, 5, 1, ["Sample", "Defects (c)", "c-bar", "UCL c", "LCL c", "Defective (np)", "np-bar",
                     "UCL np", "LCL np"])
    for i in range(100):
        r = 6 + i
        d = rows[i] if i < len(rows) else None
        c.cell(row=r, column=1, value=i + 1)
        c.cell(row=r, column=2, value=int(d["defects"]) if d else None).fill = INPUT_FILL
        c.cell(row=r, column=3, value="=AVERAGE($B$6:$B$105)")
        c.cell(row=r, column=4, value="=C%d+3*SQRT(C%d)" % (r, r))
        c.cell(row=r, column=5, value="=MAX(0,C%d-3*SQRT(C%d))" % (r, r))
        c.cell(row=r, column=6, value=int(d["defective"]) if d else None).fill = INPUT_FILL
        c.cell(row=r, column=7, value="=AVERAGE($F$6:$F$105)")
        c.cell(row=r, column=8, value="=G%d+3*SQRT(G%d*(1-G%d/$B$3))" % (r, r, r))
        c.cell(row=r, column=9, value="=MAX(0,G%d-3*SQRT(G%d*(1-G%d/$B$3)))" % (r, r, r))
    wb.save(os.path.join(HERE, "attribute-charts.xlsx"))


# ------------------------------------------------------------------ Gage RR

def gage_workbook():
    rows = read("gage_rr.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Gage R&R"
    title(ws, "Gauge R&R Study, ANOVA Method (crossed)",
          "Enter part, operator and measurement in columns A to C (balanced study: every operator "
          "measures every part the same number of times). List the part and operator IDs in F and H.")
    header(ws, 4, 1, ["Part", "Operator", "Measurement", "Cell mean"])
    first, last = 5, 504
    for i in range(first, last + 1):
        d = rows[i - first] if i - first < len(rows) else None
        ws.cell(row=i, column=1, value=int(d["part"]) if d else None).fill = INPUT_FILL
        ws.cell(row=i, column=2, value=d["operator"] if d else None).fill = INPUT_FILL
        ws.cell(row=i, column=3, value=float(d["measurement"]) if d else None).fill = INPUT_FILL
        ws.cell(row=i, column=4, value='=IF(ISNUMBER(C%d),AVERAGEIFS($C$5:$C$504,$A$5:$A$504,A%d,'
                '$B$5:$B$504,B%d),"")' % (i, i, i))
    A, B, C, D = "$A$5:$A$504", "$B$5:$B$504", "$C$5:$C$504", "$D$5:$D$504"

    header(ws, 4, 6, ["Part ID", "Part mean"])
    parts = sorted({int(r["part"]) for r in rows})
    for i in range(30):
        r = 5 + i
        ws.cell(row=r, column=6, value=parts[i] if i < len(parts) else None).fill = INPUT_FILL
        ws.cell(row=r, column=7, value='=IF(F%d="","",AVERAGEIF(%s,F%d,%s))' % (r, A, r, C))
    header(ws, 4, 8, ["Operator ID", "Operator mean"])
    ops = sorted({r["operator"] for r in rows})
    for i in range(10):
        r = 5 + i
        ws.cell(row=r, column=8, value=ops[i] if i < len(ops) else None).fill = INPUT_FILL
        ws.cell(row=r, column=9, value='=IF(H%d="","",AVERAGEIF(%s,H%d,%s))' % (r, B, r, C))

    k = 11  # results start column K
    cells = [
        ("Tolerance (USL - LSL)", 4.0, "0.000", INPUT_FILL),
        ("Study multiplier (6 = 99.73%)", 6, "0.00", INPUT_FILL),
        ("Parts (p)", '=COUNTA(F5:F34)', "0", RESULT_FILL),
        ("Operators (o)", '=COUNTA(H5:H14)', "0", RESULT_FILL),
        ("Trials (r)", "=COUNT(%s)/(L7*L8)" % C, "0", RESULT_FILL),
        ("Grand mean", "=AVERAGE(%s)" % C, "0.00000", RESULT_FILL),
        ("SS total", "=DEVSQ(%s)" % C, "0.000000", RESULT_FILL),
        ("SS part", "=L8*L9*DEVSQ(G5:G34)+L8*L9*COUNT(G5:G34)*(AVERAGE(G5:G34)-L10)^2", "0.000000", RESULT_FILL),
        ("SS operator", "=L7*L9*DEVSQ(I5:I14)+L7*L9*COUNT(I5:I14)*(AVERAGE(I5:I14)-L10)^2", "0.000000", RESULT_FILL),
        ("SS cells (part x operator means)", "=DEVSQ(%s)+COUNT(%s)*(AVERAGE(%s)-L10)^2" % (D, D, D),
         "0.000000", RESULT_FILL),
        ("SS part x operator", "=L14-L12-L13", "0.000000", RESULT_FILL),
        ("SS repeatability", "=L11-L14", "0.000000", RESULT_FILL),
    ]
    for i, (lab, f, fmt, fill) in enumerate(cells):
        label_value(ws, 5 + i, lab, f, fmt, fill, col=k)

    header(ws, 19, k, ["Source", "DF", "SS", "MS", "F", "p"])
    an = [
        ("Part", "=L7-1", "=L12", "=M20/L20", "=N20/N22", "=FDIST(O20,L20,L22)"),
        ("Operator", "=L8-1", "=L13", "=M21/L21", "=N21/N22", "=FDIST(O21,L21,L22)"),
        ("Part x Operator", "=L20*L21", "=L15", "=M22/L22", "=N22/N23", "=FDIST(O22,L22,L23)"),
        ("Repeatability", "=L7*L8*(L9-1)", "=L16", "=M23/L23", "", ""),
        ("Total", "=L7*L8*L9-1", "=L11", "", "", ""),
    ]
    for i, row in enumerate(an):
        for j, v in enumerate(row):
            c = ws.cell(row=20 + i, column=k + j, value=v)
            c.border = BOX
            if j >= 2:
                c.number_format = "0.000000" if j < 5 else "0.0000"
    label_value(ws, 26, "Pool interaction? (p > 0.25)", '=IF(P22>0.25,"Yes","No")', "@", col=k)
    label_value(ws, 27, "MS used for repeatability", '=IF(L26="Yes",(M22+M23)/(L22+L23),N23)', "0.000000", col=k)

    header(ws, 29, k, ["Source", "Variance", "% Contribution", "Std dev", "Study var", "% Study var",
                       "% Tolerance"])
    comps = [
        ("Total Gage R&R", "=L31+L32"),
        ("Repeatability", "=L27"),
        ("Reproducibility", "=L33+L34"),
        ("Operator", '=MAX(0,(N21-IF(L26="Yes",L27,N22))/(L7*L9))'),
        ("Part x Operator", '=IF(L26="Yes",0,MAX(0,(N22-N23)/L9))'),
        ("Part-to-Part", '=MAX(0,(N20-IF(L26="Yes",L27,N22))/(L8*L9))'),
        ("Total Variation", "=L30+L35"),
    ]
    for i, (lab, f) in enumerate(comps):
        r = 30 + i
        vals = [lab, f, "=100*L%d/$L$36" % r, "=SQRT(L%d)" % r, "=$L$6*N%d" % r,
                "=100*N%d/$N$36" % r, "=100*O%d/$L$5" % r]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=k + j, value=v)
            c.border = BOX
            c.number_format = "0.000000" if j in (1, 3, 4) else "0.00"
            if j and i == 0:
                c.fill = RESULT_FILL
    label_value(ws, 38, "Number of distinct categories", "=INT(1.41*N35/N30)", "0", col=k)
    label_value(ws, 39, "Verdict", '=IF(P30<10,"Acceptable",IF(P30<=30,"Marginal: depends on application",'
                '"Unacceptable: improve the gauge"))', "@", col=k)
    ws.cell(row=40, column=k, value="AIAG guidance: %Study var under 10% is acceptable, 10-30% may be "
            "acceptable, over 30% is not. ndc should be 5 or more.").font = NOTE
    widths(ws, {"A": 8, "B": 10, "C": 13, "D": 11, "F": 9, "G": 11, "H": 11, "I": 13, "K": 32,
                "L": 12, "M": 13, "N": 12, "O": 11, "P": 11, "Q": 11})
    wb.save(os.path.join(HERE, "gage-rr-anova.xlsx"))


# ---------------------------------------------------------------------- DOE

def doe_workbook():
    rows = read("doe_2k3.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "2^3 Factorial"
    title(ws, "Two-Level Full Factorial (2^3, 2 replicates)",
          "Name the factors and their low and high settings, run the experiment in random order, "
          "and enter the responses in column J.")
    header(ws, 3, 1, ["Factor", "Name", "Low (-1)", "High (+1)"])
    for i, (f, name, lo, hi) in enumerate([("A", "Temperature (C)", 150, 200), ("B", "Pressure (bar)", 2, 4),
                                           ("C", "Time (min)", 10, 20)]):
        for j, v in enumerate((f, name, lo, hi)):
            ws.cell(row=4 + i, column=1 + j, value=v).fill = INPUT_FILL if j else PatternFill()
    header(ws, 9, 1, ["Std order", "Replicate", "A", "B", "C", "AB", "AC", "BC", "ABC", "Response",
                      "A setting", "B setting", "C setting"])
    first = 10
    for i, d in enumerate(rows):
        r = first + i
        ws.cell(row=r, column=1, value=int(d["std_order"]))
        ws.cell(row=r, column=2, value=int(d["replicate"]))
        ws.cell(row=r, column=3, value=int(d["A_temperature"]))
        ws.cell(row=r, column=4, value=int(d["B_pressure"]))
        ws.cell(row=r, column=5, value=int(d["C_time"]))
        ws.cell(row=r, column=6, value="=C%d*D%d" % (r, r))
        ws.cell(row=r, column=7, value="=C%d*E%d" % (r, r))
        ws.cell(row=r, column=8, value="=D%d*E%d" % (r, r))
        ws.cell(row=r, column=9, value="=C%d*D%d*E%d" % (r, r, r))
        ws.cell(row=r, column=10, value=float(d["yield_pct"])).fill = INPUT_FILL
        for j, rowref in enumerate((4, 5, 6)):
            col = "CDE"[j]
            ws.cell(row=r, column=11 + j, value="=($C$%d+$D$%d)/2+%s%d*($D$%d-$C$%d)/2"
                    % (rowref, rowref, col, r, rowref, rowref))
    last = first + len(rows) - 1
    y = "$J$%d:$J$%d" % (first, last)
    top = last + 3
    label_value(ws, top, "Runs (N)", "=COUNT(%s)" % y, "0")
    label_value(ws, top + 1, "Grand mean", "=AVERAGE(%s)" % y, "0.0000")
    label_value(ws, top + 2, "SS total", "=DEVSQ(%s)" % y, "0.0000")
    label_value(ws, top + 3, "SS error", "=B%d-SUM(D%d:D%d)" % (top + 2, top + 7, top + 13), "0.0000")
    label_value(ws, top + 4, "DF error", "=B%d-1-7" % top, "0")
    label_value(ws, top + 5, "Std error of coefficient", "=SQRT(B%d/B%d/B%d)" % (top + 3, top + 4, top), "0.0000")
    header(ws, top + 6, 1, ["Term", "Effect", "Coefficient", "SS", "t", "p", "Significant (p<0.05)"])
    for i, (term, col) in enumerate(zip(["A", "B", "C", "AB", "AC", "BC", "ABC"], "CDEFGHI")):
        r = top + 7 + i
        rng = "$%s$%d:$%s$%d" % (col, first, col, last)
        vals = [term, "=SUMPRODUCT(%s,%s)/($B$%d/2)" % (rng, y, top), "=B%d/2" % r,
                "=SUMPRODUCT(%s,%s)^2/$B$%d" % (rng, y, top), "=C%d/$B$%d" % (r, top + 5),
                "=TDIST(ABS(E%d),$B$%d,2)" % (r, top + 4), '=IF(F%d<0.05,"Yes","")' % r]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = BOX
            if 1 <= j <= 5:
                c.number_format = "0.0000"
    ch = BarChart()
    ch.type, ch.title = "bar", "Effects (absolute size shows importance)"
    ch.add_data(Reference(ws, min_col=2, min_row=top + 6, max_row=top + 13), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=top + 7, max_row=top + 13))
    ch.height, ch.width = 8, 14
    ws.add_chart(ch, "I%d" % (top + 6))
    widths(ws, {"A": 26, "B": 16, "C": 12, "D": 12, "G": 18, "J": 11, "K": 11, "L": 11, "M": 11})
    wb.save(os.path.join(HERE, "doe-2k-factorial.xlsx"))


# --------------------------------------------------------- hypothesis tests

def hypothesis_workbook():
    rows = read("two_sample.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "t-tests"
    title(ws, "Before / After Comparison (t-tests)", "Paste up to 200 values per group. Paired tests "
          "need the same number of values in each column, in matching order.")
    header(ws, 4, 1, ["Before", "After"])
    for i in range(200):
        d = rows[i] if i < len(rows) else None
        ws.cell(row=5 + i, column=1, value=float(d["before"]) if d else None).fill = INPUT_FILL
        ws.cell(row=5 + i, column=2, value=float(d["after"]) if d else None).fill = INPUT_FILL
    a, b = "A5:A204", "B5:B204"
    k = 4
    items = [
        ("n before", "=COUNT(%s)" % a, "0"), ("n after", "=COUNT(%s)" % b, "0"),
        ("Mean before", "=AVERAGE(%s)" % a, "0.0000"), ("Mean after", "=AVERAGE(%s)" % b, "0.0000"),
        ("Std dev before", "=STDEV(%s)" % a, "0.0000"), ("Std dev after", "=STDEV(%s)" % b, "0.0000"),
        ("Difference (before - after)", "=E7-E8", "0.0000"),
        ("Welch t", "=E11/SQRT(E9^2/E5+E10^2/E6)", "0.0000"),
        ("Welch df", "=(E9^2/E5+E10^2/E6)^2/((E9^2/E5)^2/(E5-1)+(E10^2/E6)^2/(E6-1))", "0.00"),
        ("Welch p (two-sided)", "=TTEST(%s,%s,2,3)" % (a, b), "0.00000"),
        ("Pooled (equal variance) p", "=TTEST(%s,%s,2,2)" % (a, b), "0.00000"),
        ("Paired p (two-sided)", "=TTEST(%s,%s,2,1)" % (a, b), "0.00000"),
        ("F-test for equal variances p", "=FTEST(%s,%s)" % (a, b), "0.00000"),
        ("Conclusion (alpha 0.05)", '=IF(E14<0.05,"Means differ","No evidence the means differ")', "@"),
    ]
    header(ws, 4, k, ["Statistic", "Result"])
    for i, (lab, f, fmt) in enumerate(items):
        label_value(ws, 5 + i, lab, f, fmt, col=k)
    widths(ws, {"A": 11, "B": 11, "D": 32, "E": 26})
    wb.save(os.path.join(HERE, "hypothesis-tests.xlsx"))


# ---------------------------------------------------------- sigma and Pareto

def sigma_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sigma calculator"
    title(ws, "DPMO, Sigma Level and Yield Calculator", "Enter defects, units and opportunities per unit.")
    for i, (lab, v) in enumerate((("Defects found", 17), ("Units inspected", 1000),
                                  ("Opportunities per unit", 5))):
        label_value(ws, 4 + i, lab, v, "0", INPUT_FILL)
    label_value(ws, 7, "DPU (defects per unit)", "=B4/B5", "0.0000")
    label_value(ws, 8, "DPMO", "=B4/(B5*B6)*1000000", "0.0")
    label_value(ws, 9, "Sigma level (short term, with 1.5 shift)", "=NORMSINV(1-B8/1000000)+1.5", "0.00")
    label_value(ws, 10, "Long-term yield", "=1-B8/1000000", "0.0000%")

    header(ws, 12, 1, ["Step", "First-pass yield"])
    for i, v in enumerate((0.98, 0.95, 0.99, 0.97)):
        ws.cell(row=13 + i, column=1, value="Step %d" % (i + 1))
        ws.cell(row=13 + i, column=2, value=v).fill = INPUT_FILL
        ws.cell(row=13 + i, column=2).number_format = "0.00%"
    label_value(ws, 18, "Rolled throughput yield", "=PRODUCT(B13:B17)", "0.00%")

    header(ws, 20, 1, ["Sigma level", "With 1.5 shift", "DPMO", "% defective", "% yield",
                       "Cpk short term", "Cpk long term"])
    for i in range(7):
        r = 21 + i
        ws.cell(row=r, column=1, value=i + 1)
        vals = ["=A%d-1.5" % r, "=(1-NORMSDIST(B%d))*1000000" % r, "=C%d/10000" % r, "=100-D%d" % r,
                "=A%d/3" % r, "=B%d/3" % r]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + j, value=v)
            c.number_format = ["0.0", "#,##0.000", "0.00000", "0.0000000", "0.00", "0.00"][j]
    widths(ws, {"A": 40, "B": 16, "C": 14, "D": 12, "E": 14, "F": 14, "G": 14})
    wb.save(os.path.join(HERE, "sigma-level-calculator.xlsx"))


def pareto_workbook():
    rows = read("pareto.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Pareto"
    title(ws, "Pareto Analysis", "Enter categories and counts in any order (up to 30). The sorted "
          "table and chart update automatically.")
    header(ws, 4, 1, ["Category", "Count", "Sort key"])
    for i in range(30):
        r = 5 + i
        d = rows[i] if i < len(rows) else None
        ws.cell(row=r, column=1, value=d["defect"] if d else None).fill = INPUT_FILL
        ws.cell(row=r, column=2, value=int(d["count"]) if d else None).fill = INPUT_FILL
        ws.cell(row=r, column=3, value='=IF(ISNUMBER(B%d),B%d+(100-ROW())/1000000,"")' % (r, r))
    header(ws, 4, 5, ["Rank", "Category", "Count", "Percent", "Cumulative %", "Vital few (80%)"])
    for i in range(30):
        r = 5 + i
        ws.cell(row=r, column=5, value=i + 1)
        ws.cell(row=r, column=6, value='=IF(E%d<=COUNT($C$5:$C$34),INDEX($A$5:$A$34,MATCH(LARGE($C$5:$C$34,E%d),'
                '$C$5:$C$34,0)),"")' % (r, r))
        ws.cell(row=r, column=7, value='=IF(F%d="","",INDEX($B$5:$B$34,MATCH(LARGE($C$5:$C$34,E%d),'
                '$C$5:$C$34,0)))' % (r, r))
        ws.cell(row=r, column=8, value='=IF(F%d="","",G%d/SUM($B$5:$B$34))' % (r, r))
        ws.cell(row=r, column=9, value='=IF(F%d="","",SUM($G$5:G%d)/SUM($B$5:$B$34))' % (r, r))
        ws.cell(row=r, column=10, value='=IF(F%d="","",IF(I%d-H%d<0.8,"Yes",""))' % (r, r, r))
        ws.cell(row=r, column=8).number_format = ws.cell(row=r, column=9).number_format = "0.0%"
    n = len(rows)
    bar = BarChart()
    bar.title, bar.y_axis.title = "Pareto chart", "Count"
    bar.add_data(Reference(ws, min_col=7, min_row=4, max_row=4 + n), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=6, min_row=5, max_row=4 + n))
    line = LineChart()
    line.add_data(Reference(ws, min_col=9, min_row=4, max_row=4 + n), titles_from_data=True)
    line.y_axis.axId, line.y_axis.title = 200, "Cumulative %"
    line.series[0].smooth = False
    line.series[0].marker.symbol = "circle"
    line.series[0].graphicalProperties.line.solidFill = "C00000"
    line.y_axis.crosses = "max"
    bar += line
    bar.height, bar.width = 10, 20
    ws.add_chart(bar, "L4")
    widths(ws, {"A": 18, "B": 8, "C": 12, "E": 6, "F": 18, "G": 8, "H": 9, "I": 13, "J": 14})
    wb.save(os.path.join(HERE, "pareto-analysis.xlsx"))


# --------------------------------------------------------------------- FMEA

def fmea_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "PFMEA"
    title(ws, "Process FMEA Worksheet", "Rate Severity, Occurrence and Detection from 1 to 10 using the "
          "Rating Scales sheet. RPN = S x O x D. Rows with RPN of 100 or more, or Severity of 9 or 10, "
          "are highlighted.")
    cols = ["Process step / function", "Potential failure mode", "Potential effect(s)", "S",
            "Potential cause(s)", "O", "Current prevention controls", "Current detection controls", "D",
            "RPN", "S x O", "Recommended action", "Owner", "Target date", "Action taken", "New S",
            "New O", "New D", "New RPN"]
    header(ws, 4, 1, cols)
    examples = [
        ("Drill hole", "Hole oversize", "Loose fit at assembly; field leak", 7, "Worn drill bit", 5,
         "Tool life counter", "Go/no-go pin gauge, 1 in 20", 4),
        ("Drill hole", "Hole missing", "Cannot assemble; line stop", 8, "Program skipped step", 2,
         "Program verification at setup", "Vision check, 100%", 2),
        ("Deburr", "Burr remains", "Operator cut; seal damage", 9, "Brush worn", 4,
         "Brush replacement schedule", "Visual inspection", 7),
        ("Wash", "Residual oil", "Paint adhesion failure", 6, "Detergent concentration low", 3,
         "Daily titration", "Water-break test, 1 per shift", 6),
    ]
    for i in range(100):
        r = 5 + i
        ex = examples[i] if i < len(examples) else (None,) * 9
        mapping = {1: ex[0], 2: ex[1], 3: ex[2], 4: ex[3], 5: ex[4], 6: ex[5], 7: ex[6], 8: ex[7], 9: ex[8]}
        for col, v in mapping.items():
            c = ws.cell(row=r, column=col, value=v)
            c.fill, c.border = INPUT_FILL, BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=10, value='=IF(AND(ISNUMBER(D%d),ISNUMBER(F%d),ISNUMBER(I%d)),D%d*F%d*I%d,"")'
                % (r, r, r, r, r, r)).border = BOX
        ws.cell(row=r, column=11, value='=IF(AND(ISNUMBER(D%d),ISNUMBER(F%d)),D%d*F%d,"")'
                % (r, r, r, r)).border = BOX
        for col in range(12, 19):
            ws.cell(row=r, column=col).border = BOX
        ws.cell(row=r, column=19, value='=IF(AND(ISNUMBER(P%d),ISNUMBER(Q%d),ISNUMBER(R%d)),P%d*Q%d*R%d,"")'
                % (r, r, r, r, r, r)).border = BOX
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                        error="Ratings must be whole numbers from 1 to 10", showErrorMessage=True)
    ws.add_data_validation(dv)
    for col in "DFIPQR":
        dv.add("%s5:%s104" % (col, col))
    red = PatternFill("solid", fgColor="F8CBAD")
    ws.conditional_formatting.add("J5:J104", CellIsRule(operator="greaterThanOrEqual", formula=["100"], fill=red))
    ws.conditional_formatting.add("D5:D104", CellIsRule(operator="greaterThanOrEqual", formula=["9"], fill=red))
    ws.freeze_panes = "B5"
    widths(ws, {"A": 20, "B": 20, "C": 26, "D": 5, "E": 22, "F": 5, "G": 22, "H": 24, "I": 5, "J": 7,
                "K": 7, "L": 26, "M": 12, "N": 12, "O": 22, "P": 6, "Q": 6, "R": 6, "S": 8})

    sc = wb.create_sheet("Rating Scales")
    title(sc, "Rating Scales (general guidance; adapt to your industry)", "")
    header(sc, 3, 1, ["Rating", "Severity (effect)", "Occurrence (likelihood of cause)",
                      "Detection (chance controls catch it)"])
    scale = [
        (10, "Hazardous: safety or regulatory failure without warning", "Very high: failure almost certain",
         "Cannot detect, or no control"),
        (9, "Hazardous with warning", "Very high", "Very remote chance of detection"),
        (8, "Very high: product inoperable, loss of primary function", "High: repeated failures",
         "Remote"),
        (7, "High: product operable at reduced performance; customer very dissatisfied", "High",
         "Very low"),
        (6, "Moderate: comfort or convenience item inoperable", "Moderate: occasional failures", "Low"),
        (5, "Low: reduced performance of a secondary function", "Moderate", "Moderate"),
        (4, "Very low: defect noticed by most customers", "Moderate to low", "Moderately high"),
        (3, "Minor: defect noticed by average customers", "Low: relatively few failures", "High"),
        (2, "Very minor: defect noticed by discerning customers", "Low", "Very high"),
        (1, "None: no discernible effect", "Remote: failure unlikely", "Almost certain to detect"),
    ]
    for i, row in enumerate(scale):
        for j, v in enumerate(row):
            c = sc.cell(row=4 + i, column=1 + j, value=v)
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(sc, {"A": 8, "B": 50, "C": 34, "D": 34})
    os.makedirs(FMEA_DIR, exist_ok=True)
    wb.save(os.path.join(FMEA_DIR, "fmea-worksheet.xlsx"))


if __name__ == "__main__":
    for build in (capability_workbook, xbar_r_workbook, attribute_workbook, gage_workbook, doe_workbook,
                  hypothesis_workbook, sigma_workbook, pareto_workbook, fmea_workbook):
        build()
        print("built", build.__name__)
