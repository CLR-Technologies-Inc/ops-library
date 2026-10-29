"""Recalculate the workbooks in LibreOffice and compare them with the Python results.

    python check_workbooks.py

Needs LibreOffice (soffice) on the PATH and openpyxl. Exits non-zero on a mismatch.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REF = json.load(open(os.path.join(HERE, "..", "data", "reference_results.json")))

CHECKS = {
    "capability-and-imr.xlsx": [
        ("Capability", "F5", REF["capability"]["mean"]),
        ("Capability", "F6", REF["capability"]["sigma_overall"]),
        ("Capability", "F8", REF["capability"]["sigma_within"]),
        ("Capability", "F9", REF["capability"]["Cp"]),
        ("Capability", "F10", REF["capability"]["Cpk"]),
        ("Capability", "F11", REF["capability"]["Cpm"]),
        ("Capability", "F12", REF["capability"]["Pp"]),
        ("Capability", "F13", REF["capability"]["Ppk"]),
        ("Capability", "F16", REF["capability_ppm_overall"]),
        ("I-MR Chart", "D5", REF["imr"]["ucl"]),
        ("I-MR Chart", "E5", REF["imr"]["lcl"]),
        ("I-MR Chart", "H5", REF["imr"]["mr_ucl"]),
    ],
    "xbar-r-chart.xlsx": [
        ("Xbar-R", "C3", REF["xbar_r"]["xbar_cl"]),
        ("Xbar-R", "D3", REF["xbar_r"]["r_bar"]),
        ("Xbar-R", "P6", REF["xbar_r"]["xbar_ucl"]),
        ("Xbar-R", "Q6", REF["xbar_r"]["xbar_lcl"]),
        ("Xbar-R", "S6", REF["xbar_r"]["r_ucl"]),
    ],
    "attribute-charts.xlsx": [
        ("p chart", "B3", REF["p_chart"]["p_bar"]),
        ("p chart", "H%d" % (4 + REF["p_chart"]["signals"][0]), "OUT"),
    ],
    "gage-rr-anova.xlsx": [
        ("Gage R&R", "P22", REF["gage_rr"]["interaction_p"]),
        ("Gage R&R", "L31", REF["gage_rr"]["var_repeatability"]),
        ("Gage R&R", "L32", REF["gage_rr"]["var_reproducibility"]),
        ("Gage R&R", "L35", REF["gage_rr"]["var_part"]),
        ("Gage R&R", "P30", REF["gage_rr"]["pct_study_var_grr"]),
        ("Gage R&R", "L38", REF["gage_rr"]["ndc"]),
    ],
    "doe-2k-factorial.xlsx": [
        ("2^3 Factorial", "B35", REF["doe"]["A"]),
        ("2^3 Factorial", "B36", REF["doe"]["B"]),
        ("2^3 Factorial", "B37", REF["doe"]["C"]),
        ("2^3 Factorial", "B38", REF["doe"]["A*B"]),
        ("2^3 Factorial", "B41", REF["doe"]["A*B*C"]),
    ],
    "hypothesis-tests.xlsx": [
        ("t-tests", "E12", REF["welch_t"]["t"]),
        ("t-tests", "E13", REF["welch_t"]["df"]),
        ("t-tests", "E14", REF["welch_t"]["p"]),
    ],
    "sigma-level-calculator.xlsx": [
        ("Sigma calculator", "B8", 3400.0),
        ("Sigma calculator", "C26", 3.398),
    ],
    "pareto-analysis.xlsx": [
        ("Pareto", "F5", "Wrong dimension"),
        ("Pareto", "G5", 112),
        ("Pareto", "J8", "Yes"),
        ("Pareto", "J9", ""),
    ],
}


def recalc(path, outdir):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return os.path.join(outdir, os.path.basename(path))


def close(got, expected):
    if isinstance(expected, str):
        return (got or "") == expected
    if not isinstance(got, (int, float)):
        return False
    tol = 1e-3 if expected == 3.398 else 1e-6 * max(1.0, abs(expected))
    return abs(float(got) - expected) <= tol


def main():
    if not shutil.which("soffice"):
        sys.exit("LibreOffice (soffice) not found")
    failures = 0
    paths = {name: os.path.join(HERE, name) for name in CHECKS}
    paths["fmea-worksheet.xlsx"] = os.path.join(HERE, "..", "..", "..", "03-design-and-engineering", "tools",
                                                "excel", "fmea-worksheet.xlsx")
    with tempfile.TemporaryDirectory() as out:
        for name, src in paths.items():
            wb = load_workbook(recalc(src, out), data_only=True)
            # No formula may evaluate to an error
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A":
                            print("FAIL %s %s!%s evaluates to %s" % (name, ws.title, c.coordinate, c.value))
                            failures += 1
            for sheet, cell, expected in CHECKS.get(name, []):
                got = wb[sheet][cell].value
                ok = close(got, expected)
                failures += not ok
                print("%-4s %-28s %-18s %-8s got %-22r expected %r" % ("ok" if ok else "FAIL", name, sheet,
                                                                         cell, got, expected))
            if name == "fmea-worksheet.xlsx":
                rpn = wb["PFMEA"]["J5"].value
                ok = rpn == 7 * 5 * 4
                failures += not ok
                print("%-4s %-28s %-18s %-8s got %-22r expected %r" % ("ok" if ok else "FAIL", name, "PFMEA",
                                                                         "J5", rpn, 140))
    print("\n%d failure(s)" % failures)
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
